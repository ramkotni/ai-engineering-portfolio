"""Extraction: complex PDF, Word, HTML and Markdown -> normalised Markdown + metadata + extracted images.

Every format becomes Markdown so one structure-aware chunker handles all of them. The PDF path handles the
layouts that break naive text extraction (see docs/DESIGN.md section 3):

  1. Running headers / footers   lines in the top / bottom 8% of the page whose text (digits masked, so
                                  "Page 2 of 4" == "Page 3 of 4") repeats on at least half the pages are
                                  removed; the page is cropped below the header and above the footer.
  2. Tables                       pdfplumber finds ruled tables; each becomes a Markdown pipe table.
  3. Multi-page tables            a table whose header row equals the previous table's header row is a
                                  continuation: its rows are appended and the repeated header dropped;
                                  "Table N (continued)" captions are removed.
  4. Images                       PyMuPDF extracts embedded images with their position on the page; each is
                                  saved as PNG and replaced in the text by [[IMAGE:<id>]] at that position,
                                  so its vision caption lands under the right section heading.
  5. Reading order / columns      the page is cut into horizontal strips around tables and images; inside a
                                  strip, each text row is full-width if any character crosses the page's
                                  centre gutter, otherwise it belongs to a column. Runs of column rows are read
                                  left column first, then right column.
  6. Headings                     font size >= 1.15 x the median body size; the three largest sizes map to #, ##, ###.
  7. Footnotes                    text smaller than 0.8 x body size is kept as "Footnote: ..." lines.
  8. Metadata                     "Document ID: X | Revision: 3 | Classification: Public" lines -> metadata.
"""

import math
import re
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup

IMAGE_MARKER = re.compile(r"\[\[IMAGE:([\w.-]+)\]\]")
_META_KEYS = {"document id": "doc_id", "revision": "revision", "owner": "owner", "classification": "classification"}
MIN_IMAGE_PIXELS = 120 * 120     # ignore logos / bullets
HF_BAND = 0.08                   # header / footer band = top and bottom 8% of the page
GUTTER_HALF_WIDTH = 4            # a row is full-width if a character overlaps centre +/- 4 pt
CONTINUED = re.compile(r"^(table|figure)\s+\d+\s*\(continued\)$", re.IGNORECASE)
CAPTION = re.compile(r"^(table|figure)\s*\d+[.:]", re.IGNORECASE)


@dataclass
class LoadedDocument:
    path: Path
    markdown: str
    meta: dict = field(default_factory=dict)
    images: list[dict] = field(default_factory=list)   # {"image_id", "path", "page", "caption_hint"}
    stats: dict = field(default_factory=dict)          # what the extractor did - shown in the UI and DESIGN.md

    @property
    def doc_id(self) -> str:
        return str(self.meta.get("doc_id") or self.path.stem.split("_")[0])

    @property
    def title(self) -> str:
        return str(self.meta.get("title") or self.path.stem)

    @property
    def classification(self) -> str:
        return str(self.meta.get("classification", "internal")).lower()


def parse_meta_line(text: str, meta: dict) -> bool:
    """'Document ID: X | Revision: 3 | Classification: Public' -> meta. Returns True if it was a meta line."""
    found = False
    for part in text.split("|"):
        key, _, value = part.partition(":")
        if key.strip().lower() in _META_KEYS and value.strip():
            meta[_META_KEYS[key.strip().lower()]] = value.strip()
            found = True
    return found


def table_md(rows: list[list]) -> str:
    rows = [[" ".join(str(c or "").split()).replace("|", "/") for c in r] for r in rows if r and any(r)]
    if not rows:
        return ""
    width = max(map(len, rows))
    rows = [r + [""] * (width - len(r)) for r in rows]
    return "\n".join(["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width] + ["| " + " | ".join(r) + " |" for r in rows[1:]])


# ---------------------------------------------------------------- PDF
@dataclass
class _Line:
    text: str
    size: float
    bold: bool


def _hf_key(text: str) -> str:
    return re.sub(r"\d+", "#", " ".join(text.split())).lower()


def running_lines(pdf) -> set[str]:
    """Header/footer texts: in the top/bottom band and repeated (digits masked) on >= half of the pages."""
    counts: Counter = Counter()
    for page in pdf.pages:
        band = {_hf_key(ln["text"]) for ln in page.extract_text_lines()
                if ln["top"] < page.height * HF_BAND or ln["bottom"] > page.height * (1 - HF_BAND)}
        counts.update(band)
    n = len(pdf.pages)
    return {k for k, c in counts.items() if n >= 2 and c >= max(2, math.ceil(n / 2))}


def _body_bounds(page, running: set[str]) -> tuple[float, float, int]:
    """Crop line just below the running header and just above the running footer."""
    top, bottom, removed = 0.0, float(page.height), 0
    for ln in page.extract_text_lines():
        if _hf_key(ln["text"]) not in running:
            continue
        removed += 1
        if ln["top"] < page.height * HF_BAND:
            top = max(top, ln["bottom"] + 1)
        else:
            bottom = min(bottom, ln["top"] - 1)
    return top, bottom, removed


def _lines(region) -> list[_Line]:
    out = []
    for line in region.extract_text_lines(return_chars=True):
        text = " ".join(line["text"].split())
        if text:
            chars = [c for c in line["chars"] if c["text"].strip()]
            out.append(_Line(text, max(c["size"] for c in chars), all("bold" in c["fontname"].lower() for c in chars)))
    return out


def _text_in_reading_order(region, stats: Counter) -> list[_Line]:
    """Full-width rows read top to bottom; runs of two-column rows read left column, then right column."""
    x0, top, x1, bottom = region.bbox
    if bottom - top < 1:
        return []
    mid = (x0 + x1) / 2
    rows: dict[int, list[dict]] = defaultdict(list)
    for c in region.chars:
        rows[round(c["top"])].append(c)
    if not rows:
        return []
    sizes = [c["size"] for c in region.chars if c["text"].strip()]
    region_body = statistics.median(sizes) if sizes else 0
    segments: list[list] = []          # [is_full_width, y_top, y_bottom]
    for key in sorted(rows):
        row = rows[key]
        text = "".join(c["text"] for c in sorted(row, key=lambda c: c["x0"])).strip()
        crosses = any(c["x0"] < mid + GUTTER_HALF_WIDTH and c["x1"] > mid - GUTTER_HALF_WIDTH for c in row)
        both_sides = any(c["x1"] <= mid for c in row) and any(c["x0"] >= mid for c in row)
        # captions and headings break the column flow even when they are too short to cross the gutter -
        # unless the row has text on both sides (two column headings side by side)
        full = crosses or (not both_sides and (bool(CAPTION.match(text)) or max(c["size"] for c in row) >= region_body * 1.15))
        y0, y1 = min(c["top"] for c in row), max(c["bottom"] for c in row)
        if segments and segments[-1][0] == full:
            segments[-1][2] = max(segments[-1][2], y1)
        else:
            segments.append([full, y0, y1])
    out: list[_Line] = []
    for full, y0, y1 in segments:
        y0, y1 = max(top, y0 - 0.5), min(bottom, y1 + 0.5)
        if full:
            out += _lines(region.crop((x0, y0, x1, y1)))
        else:
            left, right = _lines(region.crop((x0, y0, mid, y1))), _lines(region.crop((mid, y0, x1, y1)))
            if left and right:
                stats["two_column_blocks"] += 1
            out += left + right
    return out


def load_pdf(path: Path, image_dir: Path) -> LoadedDocument:
    import pdfplumber
    import pymupdf

    meta: dict = {}
    images: list[dict] = []
    stats: Counter = Counter()
    segments: list[tuple[str, object]] = []
    prefix = path.stem.split("_")[0]
    mu = pymupdf.open(str(path))
    with pdfplumber.open(str(path)) as pdf:
        if (pdf.metadata or {}).get("Title"):
            meta["title"] = pdf.metadata["Title"]
        running = running_lines(pdf)
        stats["pages"] = len(pdf.pages)
        for pno, page in enumerate(pdf.pages):
            body_top, body_bottom, removed = _body_bounds(page, running)
            stats["header_footer_lines_removed"] += removed
            body = page.crop((0, body_top, page.width, body_bottom))
            blocks = [("table", t.bbox[1], t.bbox[3], t) for t in body.find_tables()]
            for n, info in enumerate(mu[pno].get_images(full=True), start=1):
                xref = info[0]
                img = mu.extract_image(xref)
                if img["width"] * img["height"] < MIN_IMAGE_PIXELS:
                    continue
                rects = mu[pno].get_image_rects(xref)
                y0, y1 = (rects[0].y0, rects[0].y1) if rects else (body_bottom, body_bottom)
                image_id = f"{prefix}_p{pno + 1}_img{n}"
                out = image_dir / f"{image_id}.png"
                pymupdf.Pixmap(mu, xref).save(str(out)) if img["ext"] != "png" else out.write_bytes(img["image"])
                images.append({"image_id": image_id, "path": str(out), "page": pno + 1})
                blocks.append(("image", max(body_top, y0), min(body_bottom, y1), image_id))
            y = body_top
            for kind, b_top, b_bottom, obj in sorted(blocks, key=lambda b: b[1]):
                segments.append(("lines", _text_in_reading_order(body.crop((0, y, page.width, max(y, b_top))), stats)))
                segments.append(("table", obj.extract()) if kind == "table" else ("image", obj))
                y = max(y, b_bottom)
            segments.append(("lines", _text_in_reading_order(body.crop((0, min(y, body_bottom), page.width, body_bottom)), stats)))

    sizes = [ln.size for kind, lines in segments if kind == "lines" for ln in lines]
    body_size = statistics.median(sizes) if sizes else 0
    heading_sizes = sorted({round(s, 1) for s in sizes if s >= body_size * 1.15}, reverse=True)[:3]
    parts: list[str] = []
    last_table_header: list | None = None
    for kind, content in segments:
        if kind == "table":
            rows = [r for r in content if r and any(r)]
            stats["tables_found"] += 1
            if parts and parts[-1].startswith("|") and last_table_header == rows[0]:   # continuation of the previous table
                parts[-1] += "\n" + "\n".join(table_md([rows[0], r]).splitlines()[2] for r in rows[1:])
                stats["tables_merged_across_pages"] += 1
            else:
                parts.append(table_md(rows))
                last_table_header = rows[0]
            continue
        if kind == "image":
            parts.append(f"[[IMAGE:{content}]]")
            continue
        for ln in content:
            if CONTINUED.match(ln.text):
                continue
            if round(ln.size, 1) in heading_sizes:
                level = heading_sizes.index(round(ln.size, 1)) + 1
                if parts and parts[-1].startswith("#" * level + " ") and level == 1:   # a title wrapped over two lines
                    parts[-1] += " " + ln.text
                else:
                    parts.append(f"{'#' * level} {ln.text}")
                if level == 1:
                    meta.setdefault("title", ln.text)
            elif parse_meta_line(ln.text, meta):
                continue
            elif ln.size < body_size * 0.8:
                parts.append(f"Footnote: {ln.text}")
                stats["footnotes"] += 1
            elif parts and not parts[-1].startswith(("#", "|", "[[", "Footnote:")) and not parts[-1].rstrip().endswith((".", ":", ";", "?", "!")):
                parts[-1] += " " + ln.text
            else:
                parts.append(ln.text)
    _attach_figure_captions(parts, images)
    stats["images_extracted"] = len(images)
    return LoadedDocument(path, "\n\n".join(p for p in parts if p), meta, images, dict(stats))


def _attach_figure_captions(parts: list[str], images: list[dict]) -> None:
    """The printed caption ("Figure 1. ...") right after an image marker becomes the image's caption hint."""
    for i, part in enumerate(parts):
        m = IMAGE_MARKER.fullmatch(part)
        if m and i + 1 < len(parts) and re.match(r"^figure\s+\d+", parts[i + 1], re.IGNORECASE):
            for img in images:
                if img["image_id"] == m.group(1):
                    img["caption_hint"] = parts[i + 1]


# ---------------------------------------------------------------- Word
def load_docx(path: Path, image_dir: Path) -> LoadedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    meta: dict = {"title": d.core_properties.title} if d.core_properties.title else {}
    parts: list[str] = []
    images: list[dict] = []
    stats: Counter = Counter()
    prefix = path.stem.split("_")[0]
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            t = Table(child, d)
            parts.append(table_md([[c.text for c in row.cells] for row in t.rows]))
            stats["tables_found"] += 1
            continue
        if tag != "p":
            continue
        p = Paragraph(child, d)
        for blip in child.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            part = d.part.related_parts[rid]
            image_id = f"{prefix}_img{len(images) + 1}"
            out = image_dir / f"{image_id}.{part.partname.ext}"
            out.write_bytes(part.blob)
            images.append({"image_id": image_id, "path": str(out), "page": None})
            parts.append(f"[[IMAGE:{image_id}]]")
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style == "title" or style.startswith("heading"):
            level = 1 if style == "title" else min(int(re.sub(r"\D", "", style) or 1) + 1, 3)
            parts.append(f"{'#' * level} {text}")
            if level == 1:
                meta.setdefault("title", text)
        elif style.startswith("list"):
            parts.append(f"- {text}")
        elif not parse_meta_line(text, meta):
            parts.append(text)
    stats["images_extracted"] = len(images)
    return LoadedDocument(path, "\n\n".join(parts), meta, images, dict(stats))


# ---------------------------------------------------------------- HTML + Markdown
def load_html(path: Path) -> LoadedDocument:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    meta = {"title": soup.title.get_text(strip=True)} if soup.title else {}
    parts: list[str] = []
    stats: Counter = Counter()
    for el in soup.body.find_all(["h1", "h2", "h3", "p", "ul", "ol", "table"]):
        if el.name not in ("ul", "ol", "table") and el.find_parent(["table", "ul", "ol"]):
            continue
        text = " ".join(el.get_text(" ", strip=True).split())
        if el.name in ("h1", "h2", "h3"):
            parts.append("#" * int(el.name[1]) + " " + text)
        elif el.name == "table":
            parts.append(table_md([[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])] for tr in el.find_all("tr")]))
            stats["tables_found"] += 1
        elif el.name in ("ul", "ol"):
            parts.append("\n".join(f"- {' '.join(li.get_text(' ', strip=True).split())}" for li in el.find_all("li")))
        elif text and not parse_meta_line(text, meta):
            parts.append(text)
    return LoadedDocument(path, "\n\n".join(parts), meta, stats=dict(stats))


def load_markdown(path: Path) -> LoadedDocument:
    meta: dict = {}
    lines = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if parse_meta_line(line, meta):
            continue
        if line.startswith("# ") and "title" not in meta:
            meta["title"] = line[2:].strip()
        lines.append(line)
    return LoadedDocument(path, "\n".join(lines), meta)


def load_document(path: Path, image_dir: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path, image_dir)
    if suffix == ".docx":
        return load_docx(path, image_dir)
    if suffix in (".html", ".htm"):
        return load_html(path)
    if suffix in (".md", ".txt"):
        return load_markdown(path)
    raise ValueError(f"Unsupported: {path}")


if __name__ == "__main__":   # python -m medbrief.loaders data/raw/<file>  -> print the extraction
    import sys
    import tempfile

    doc = load_document(Path(sys.argv[1]), Path(tempfile.mkdtemp()))
    print(doc.meta, doc.stats, [i["image_id"] for i in doc.images], sep="\n")
    print("-" * 80)
    print(doc.markdown)
