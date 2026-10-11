"""Generates the synthetic source set for the Condition Briefing Assistant.

    python data/generate_data.py

Outputs (committed) in data/raw/:
  CARD-GL-2026_hfref_guideline_summary.pdf   COMPLEX PDF (public) - running header/footer on every page, two-column body
                                             with full-width headings, a ruled table that continues across a page break
                                             (header row repeated), a treatment-algorithm flowchart and a pipeline bar
                                             chart (content only in the images), small-font footnotes
  DM-GL-2026_type2_diabetes_care_summary.html HTML guideline summary with a drug-class table (public)
  STRAT-2026_cardiometabolic_service_line.docx Word strategy brief with a table and an embedded chart (internal)
  pipeline_trials.csv                         synthetic trial registry: 26 programs, sponsors, lead institutions (public)

Clinical content is a simplified, SYNTHETIC summary for demonstrating document AI - not clinical guidance.
All companies, institutions, trial IDs, products and health-system figures are fictional.
"""

import csv
import io
from collections import Counter
from pathlib import Path

import docx
import pymupdf
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

RAW = Path(__file__).parent / "raw"
DISCLAIMER = "SYNTHETIC DATA - NOT FOR CLINICAL USE"

# ---------------------------------------------------------------- trial registry (also drives Figure 2)
TRIALS = [
    # trial_id, condition, intervention, mechanism, modality, phase, status, sponsor, sponsor_type, lead_institution, enrollment, start, primary_completion, endpoint
    ("SYN-HF-101", "HFrEF", "CRD-417 (oral cardiac myosin activator)", "Cardiac myosin activation", "Small molecule", "Phase 3", "Recruiting",
     "Cardiora Therapeutics", "Industry", "Lakeshore University Heart Institute", 2400, "2025-03", "2028-06", "Time to CV death or first HF event"),
    ("SYN-HF-102", "HFrEF", "CRD-522 (next-generation myosin activator)", "Cardiac myosin activation", "Small molecule", "Phase 2", "Recruiting",
     "Cardiora Therapeutics", "Industry", "Pacific Crest Medical Center", 320, "2026-01", "2027-09", "Change in LVEF at 24 weeks"),
    ("SYN-HF-103", "HFrEF", "HLX-SERCA (AAV SERCA2a gene therapy)", "SERCA2a gene transfer", "Gene therapy", "Phase 2", "Active, not recruiting",
     "Helix Biologics", "Industry", "Northbridge Clinical Research Network", 150, "2024-09", "2026-12", "Peak VO2 at 12 months"),
    ("SYN-HF-104", "HFrEF", "NRD-88 (nonsteroidal MRA)", "Nonsteroidal mineralocorticoid receptor antagonism", "Small molecule", "Phase 3", "Recruiting",
     "Nordlicht Pharma", "Industry", "Lakeshore University Heart Institute", 3100, "2025-06", "2029-03", "CV death or HF hospitalization"),
    ("SYN-HF-105", "HFrEF", "MER-Sense (implantable pulmonary artery pressure sensor)", "Remote hemodynamic monitoring", "Device", "Pivotal", "Recruiting",
     "Meridian Cardio Devices", "Industry", "Pacific Crest Medical Center", 1000, "2025-02", "2027-12", "HF hospitalizations at 12 months"),
    ("SYN-HF-106", "HFrEF", "MER-Flow (interatrial shunt)", "Interatrial shunt", "Device", "Phase 2", "Recruiting",
     "Meridian Cardio Devices", "Industry", "Northbridge Clinical Research Network", 240, "2025-10", "2027-10", "Composite of death, HF events, KCCQ"),
    ("SYN-HF-107", "HFrEF", "TitrateAI (AI-guided GDMT titration program)", "Digital GDMT optimization", "Digital", "Pragmatic RCT", "Recruiting",
     "Pentara Health AI", "Industry", "Summit Valley Health Partners", 5000, "2025-09", "2027-03", "Share of patients on all four GDMT pillars at 6 months"),
    ("SYN-HF-108", "HFrEF", "Pharmacist-led rapid titration clinic", "Care-delivery redesign", "Care model", "Pragmatic RCT", "Recruiting",
     "Lakeshore University Heart Institute", "Academic", "Lakeshore University Heart Institute", 600, "2025-04", "2026-11", "GDMT target-dose attainment at 90 days"),
    ("SYN-HF-109", "HFrEF", "HLX-207 (anti-fibrotic monoclonal antibody)", "Myocardial fibrosis inhibition", "Biologic", "Phase 1", "Recruiting",
     "Helix Biologics", "Industry", "Northbridge Clinical Research Network", 60, "2026-02", "2027-02", "Safety and tolerability"),
    ("SYN-HF-110", "HFrEF", "CRD-417 plus SGLT2 inhibitor", "Cardiac myosin activation", "Small molecule", "Phase 2", "Completed",
     "Cardiora Therapeutics", "Industry", "Lakeshore University Heart Institute", 280, "2023-04", "2025-08", "NT-proBNP change at 16 weeks"),
    ("SYN-HF-111", "HFrEF", "NRD-31 (oral sGC stimulator)", "Soluble guanylate cyclase stimulation", "Small molecule", "Phase 2", "Recruiting",
     "Nordlicht Pharma", "Industry", "Pacific Crest Medical Center", 400, "2025-11", "2027-08", "NT-proBNP change at 12 weeks"),
    ("SYN-HF-112", "HFrEF", "MER-CCM2 (cardiac contractility modulation)", "Device-based contractility modulation", "Device", "Phase 1", "Not yet recruiting",
     "Meridian Cardio Devices", "Industry", "Pacific Crest Medical Center", 40, "2026-09", "2027-12", "Device safety at 6 months"),
    ("SYN-HF-113", "HFrEF", "Hospital-to-home IV iron repletion pathway", "Iron deficiency correction", "Care model", "Phase 4", "Recruiting",
     "Northbridge Clinical Research Network", "Academic", "Northbridge Clinical Research Network", 900, "2025-01", "2026-10", "6-minute walk distance at 24 weeks"),
    ("SYN-HF-114", "HFrEF", "HLX-SERCA long-term follow-up", "SERCA2a gene transfer", "Gene therapy", "Phase 1/2", "Active, not recruiting",
     "Helix Biologics", "Industry", "Northbridge Clinical Research Network", 45, "2023-06", "2028-06", "Long-term safety at 5 years"),
    ("SYN-DM-201", "Type 2 diabetes", "GLV-12 (oral small-molecule GLP-1 receptor agonist)", "GLP-1 receptor agonism (oral)", "Small molecule", "Phase 3", "Recruiting",
     "Glycova Bio", "Industry", "Summit Valley Diabetes Center", 2800, "2025-05", "2027-11", "A1C change at 40 weeks"),
    ("SYN-DM-202", "Type 2 diabetes", "GLV-30 (weekly GIP/GLP-1/glucagon triple agonist)", "Triple incretin agonism", "Biologic", "Phase 2", "Recruiting",
     "Glycova Bio", "Industry", "Pacific Crest Medical Center", 420, "2025-12", "2027-06", "Body-weight change at 36 weeks"),
    ("SYN-DM-203", "Type 2 diabetes", "HLX-INS7 (once-weekly basal insulin)", "Once-weekly basal insulin", "Biologic", "Phase 3", "Active, not recruiting",
     "Helix Biologics", "Industry", "Summit Valley Diabetes Center", 1500, "2024-08", "2026-09", "A1C change at 52 weeks"),
    ("SYN-DM-204", "Type 2 diabetes", "NRD-560 (hepatoselective glucokinase activator)", "Glucokinase activation", "Small molecule", "Phase 2", "Recruiting",
     "Nordlicht Pharma", "Industry", "Lakeshore University Heart Institute", 360, "2026-03", "2027-09", "A1C change at 24 weeks"),
    ("SYN-DM-205", "Type 2 diabetes", "CGM-linked digital coaching (Pentara Glide)", "Digital lifestyle and CGM coaching", "Digital", "Pragmatic RCT", "Recruiting",
     "Pentara Health AI", "Industry", "Summit Valley Health Partners", 3000, "2025-07", "2026-12", "A1C below 7% at 12 months"),
    ("SYN-DM-206", "Type 2 diabetes", "Encapsulated stem-cell-derived islets", "Beta-cell replacement", "Cell therapy", "Phase 1/2", "Recruiting",
     "Summit Valley Diabetes Center", "Academic", "Summit Valley Diabetes Center", 30, "2025-10", "2028-04", "Insulin independence at 12 months"),
    ("SYN-DM-207", "Type 2 diabetes", "Structured low-energy diet remission program", "Diabetes remission through weight loss", "Care model", "Pragmatic RCT", "Recruiting",
     "Lakeshore University Heart Institute", "Academic", "Lakeshore University Heart Institute", 800, "2025-02", "2027-02", "Remission (A1C below 6.5% off medication) at 12 months"),
    ("SYN-DM-208", "Type 2 diabetes", "GLV-12 cardiovascular outcomes study", "GLP-1 receptor agonism (oral)", "Small molecule", "Phase 3", "Recruiting",
     "Glycova Bio", "Industry", "Northbridge Clinical Research Network", 9000, "2025-09", "2030-06", "Major adverse cardiovascular events"),
    ("SYN-DM-209", "Type 2 diabetes", "MER-Gluco 180-day implantable CGM", "Long-term continuous glucose monitoring", "Device", "Pivotal", "Recruiting",
     "Meridian Cardio Devices", "Industry", "Pacific Crest Medical Center", 350, "2025-11", "2027-05", "Sensor accuracy (MARD) over 180 days"),
    ("SYN-DM-210", "Type 2 diabetes", "HLX-AMY (amylin analog co-formulation)", "Amylin and GLP-1 co-agonism", "Biologic", "Phase 1", "Recruiting",
     "Helix Biologics", "Industry", "Northbridge Clinical Research Network", 80, "2026-04", "2027-04", "Safety and pharmacokinetics"),
    ("SYN-DM-211", "Type 2 diabetes", "NRD-560 plus metformin", "Glucokinase activation", "Small molecule", "Phase 2", "Completed",
     "Nordlicht Pharma", "Industry", "Pacific Crest Medical Center", 240, "2023-09", "2025-06", "A1C change at 12 weeks"),
    ("SYN-DM-212", "Type 2 diabetes", "Community pharmacist insulin titration", "Care-delivery redesign", "Care model", "Phase 4", "Recruiting",
     "Northbridge Clinical Research Network", "Academic", "Northbridge Clinical Research Network", 700, "2025-06", "2026-12", "Time in range at 6 months"),
]
FIELDS = ["trial_id", "condition", "intervention", "mechanism", "modality", "phase", "status", "sponsor", "sponsor_type",
          "lead_institution", "enrollment", "start_date", "primary_completion", "primary_endpoint"]


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    for name in (("arialbd.ttf", "DejaVuSans-Bold.ttf") if bold else ("arial.ttf", "DejaVuSans.ttf")):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def _png(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- figures (content exists ONLY in the image)
def algorithm_png() -> bytes:
    img = Image.new("RGB", (1200, 820), "white")
    d = ImageDraw.Draw(img)
    f, fb = font(22), font(26, bold=True)
    d.text((40, 18), "HFrEF treatment algorithm (synthetic summary)", fill="black", font=fb)

    def box(x, y, w, h, text, fill):
        d.rounded_rectangle((x, y, x + w, y + h), radius=14, outline="black", width=3, fill=fill)
        for i, line in enumerate(text.split("\n")):
            d.text((x + 16, y + 12 + i * 28), line, fill="black", font=f)

    def down(x, y1, y2):
        d.line((x, y1, x, y2), fill="black", width=3)
        d.polygon([(x, y2), (x - 9, y2 - 15), (x + 9, y2 - 15)], fill="black")

    box(330, 70, 540, 80, "Stage C HFrEF: LVEF 40% or lower,\nsymptoms present or prior", "#FFF2CC")
    down(600, 150, 200)
    box(250, 200, 700, 110, "Start all four pillars within 6 weeks:\nARNI (or ACEi/ARB) + beta-blocker + MRA + SGLT2i\nadd loop diuretic for congestion", "#DDEBF7")
    down(600, 310, 360)
    box(330, 360, 540, 80, "Reassess LVEF after 3 months\nof optimized GDMT", "#DDEBF7")
    d.line((600, 440, 600, 470), fill="black", width=3)
    d.line((170, 470, 1030, 470), fill="black", width=3)
    for x in (170, 600, 1030):
        down(x, 470, 520)
    box(20, 520, 330, 140, "LVEF 35% or lower,\nNYHA II-III,\nlife expectancy >1 year:\nICD", "#E2F0D9")
    box(400, 520, 400, 140, "LVEF 35% or lower,\nLBBB with QRS 150 ms\nor longer, sinus rhythm:\nCRT-D", "#E2F0D9")
    box(850, 520, 330, 140, "Persistent NYHA IV\ndespite GDMT (Stage D):\nrefer for LVAD or\ntransplant evaluation", "#F8CBAD")
    d.text((40, 700), "Also: SGLT2i is started regardless of diabetes status; titrate to target doses every 2 weeks as tolerated.", fill="black", font=f)
    d.text((40, 740), "Hydralazine + isosorbide dinitrate: add for self-identified Black patients with NYHA III-IV on optimal GDMT.", fill="black", font=f)
    return _png(img)


def pipeline_chart_png(condition: str) -> bytes:
    """Bar chart of synthetic pipeline programs by modality and phase - numbers appear only in the chart."""
    rows = [t for t in TRIALS if t[1] == condition and t[6] != "Completed"]
    counts = Counter(t[4] for t in rows)
    late = Counter(t[4] for t in rows if t[5] in ("Phase 3", "Pivotal", "Pragmatic RCT", "Phase 4"))
    img = Image.new("RGB", (1100, 600), "white")
    d = ImageDraw.Draw(img)
    f, fb = font(21), font(25, bold=True)
    d.text((30, 15), f"Active {condition} pipeline by modality (synthetic registry, 2026)", fill="black", font=fb)
    x0, y0, h = 260, 80, 420
    bar_h = 46
    peak = max(counts.values())
    for i, (modality, n) in enumerate(sorted(counts.items(), key=lambda t: -t[1])):
        y = y0 + i * (bar_h + 18)
        d.text((30, y + 10), modality, fill="black", font=f)
        w = int(n / peak * 520)
        d.rectangle((x0, y, x0 + w, y + bar_h), fill="#9DC3E6", outline="black")
        lw = int(late[modality] / peak * 520)
        if lw:
            d.rectangle((x0, y, x0 + lw, y + bar_h), fill="#1F4E79", outline="black")
        d.text((x0 + w + 12, y + 10), f"{n} programs ({late[modality]} late-stage)", fill="black", font=f)
    d.rectangle((30, y0 + h + 40, 60, y0 + h + 62), fill="#1F4E79")
    d.text((70, y0 + h + 38), "late-stage (Phase 3, pivotal, pragmatic RCT, Phase 4)", fill="black", font=f)
    d.rectangle((640, y0 + h + 40, 670, y0 + h + 62), fill="#9DC3E6")
    d.text((680, y0 + h + 38), "earlier stage", fill="black", font=f)
    return _png(img)


def admissions_chart_png() -> bytes:
    img = Image.new("RGB", (1000, 520), "white")
    d = ImageDraw.Draw(img)
    f, fb = font(20), font(24, bold=True)
    d.text((30, 15), "Riverbend Health HF admissions and 30-day readmission rate by quarter", fill="black", font=fb)
    quarters = ["Q3-24", "Q4-24", "Q1-25", "Q2-25", "Q3-25", "Q4-25", "Q1-26", "Q2-26"]
    admissions = [498, 541, 590, 552, 515, 548, 601, 566]
    readmit = [23.9, 23.4, 24.8, 22.9, 22.1, 21.8, 22.6, 21.4]
    x0, y0, w, h = 90, 70, 840, 340
    d.rectangle((x0, y0, x0 + w, y0 + h), outline="black", width=2)
    bw = w / len(quarters)
    for i, (q, a, r) in enumerate(zip(quarters, admissions, readmit)):
        bx = x0 + i * bw + 15
        top = y0 + h - a / 650 * h
        d.rectangle((bx, top, bx + bw - 30, y0 + h), fill="#9DC3E6", outline="black")
        d.text((bx, top - 26), str(a), fill="black", font=f)
        d.text((bx, y0 + h + 8), q, fill="black", font=f)
    pts = [(x0 + i * bw + bw / 2, y0 + h - (r - 18) / 10 * h) for i, r in enumerate(readmit)]
    d.line(pts, fill="#C00000", width=4)
    d.text((pts[-1][0] - 120, pts[-1][1] - 34), "readmission 21.4%", fill="#C00000", font=f)
    d.text((30, 460), "Bars: HF admissions (count). Line: 30-day all-cause readmission rate (%). Peak admissions Q1-26.", fill="black", font=f)
    return _png(img)


# ---------------------------------------------------------------- PDF layout engine (PyMuPDF)
class PdfWriter:
    W, H = 612, 792
    LEFT, RIGHT, TOP, BOTTOM = 50, 562, 62, 735
    GUTTER = 22

    def __init__(self, title: str, header: str):
        self.doc = pymupdf.open()
        self.doc.set_metadata({"title": title, "author": "Synthetic Guideline Consortium"})
        self.header = header
        self.page = None
        self.y = 0.0
        self.new_page()

    def new_page(self):
        self.page = self.doc.new_page(width=self.W, height=self.H)
        self.y = self.TOP

    @staticmethod
    def wrap(text: str, size: float, width: float, fontname: str = "helv") -> list[str]:
        lines, cur = [], ""
        for word in text.split():
            trial = f"{cur} {word}".strip()
            if pymupdf.get_text_length(trial, fontname=fontname, fontsize=size) <= width:
                cur = trial
            else:
                if cur:
                    lines.append(cur)
                cur = word
        return lines + ([cur] if cur else [])

    def _ensure(self, height: float):
        if self.y + height > self.BOTTOM:
            self.new_page()

    def heading(self, text: str, size: float = 13, keep_with_next: float = 160):
        self._ensure(size * 3 + keep_with_next)   # never strand a heading at the bottom of a page
        self.y += 6
        self.page.insert_text((self.LEFT, self.y + size), text, fontname="hebo", fontsize=size)
        self.y += size * 1.7

    def para(self, text: str, size: float = 10, fontname: str = "helv", gap: float = 6):
        lh = size * 1.3
        for line in self.wrap(text, size, self.RIGHT - self.LEFT, fontname):
            self._ensure(lh)
            self.page.insert_text((self.LEFT, self.y + size), line, fontname=fontname, fontsize=size)
            self.y += lh
        self.y += gap

    def columns(self, left: list[tuple[str, str]], right: list[tuple[str, str]], size: float = 9.5):
        """Two-column block: items are ("h", heading) or ("p", paragraph). Both columns start at the same y."""
        col_w = (self.RIGHT - self.LEFT - self.GUTTER) / 2
        lh = size * 1.3

        def height(items):
            return sum(4 + 11.5 * 1.6 if kind == "h" else len(self.wrap(text, size, col_w)) * lh + 5 for kind, text in items)

        self._ensure(max(height(left), height(right)) + 6)
        bottoms = []
        for x, items in ((self.LEFT, left), (self.LEFT + col_w + self.GUTTER, right)):
            y = self.y
            for kind, text in items:
                if kind == "h":
                    y += 4
                    self.page.insert_text((x, y + 11.5), text, fontname="hebo", fontsize=11.5)
                    y += 11.5 * 1.6
                    continue
                for line in self.wrap(text, size, col_w):
                    self.page.insert_text((x, y + size), line, fontname="helv", fontsize=size)
                    y += lh
                y += 5
            bottoms.append(y)
        self.y = max(bottoms) + 6

    def table(self, title: str, header: list[str], rows: list[list[str]], widths: list[float], size: float = 8.2):
        lh, pad = size * 1.25, 4

        def row_height(cells):
            return max(len(self.wrap(c, size, w - 2 * pad)) for c, w in zip(cells, widths)) * lh + 2 * pad

        def draw_row(cells, bold=False):
            h = row_height(cells)
            x = self.LEFT
            for c, w in zip(cells, widths):
                self.page.draw_rect(pymupdf.Rect(x, self.y, x + w, self.y + h), color=(0, 0, 0), width=0.6,
                                    fill=(0.87, 0.92, 0.97) if bold else None)
                for i, line in enumerate(self.wrap(c, size, w - 2 * pad)):
                    self.page.insert_text((x + pad, self.y + pad + size + i * lh), line, fontname="hebo" if bold else "helv", fontsize=size)
                x += w
            self.y += h

        self._ensure(row_height(header) + row_height(rows[0]) + 30)
        self.page.insert_text((self.LEFT, self.y + 10), title, fontname="hebo", fontsize=10)
        self.y += 16
        draw_row(header, bold=True)
        for r in rows:
            if self.y + row_height(r) > self.BOTTOM:          # continue on the next page, repeating the header row
                self.new_page()
                self.page.insert_text((self.LEFT, self.y + 10), f"{title.split('.')[0]} (continued)", fontname="hebo", fontsize=10)
                self.y += 16
                draw_row(header, bold=True)
            draw_row(r)
        self.y += 10

    def figure(self, png: bytes, height: float, caption: str):
        self._ensure(height + 30)
        width = self.RIGHT - self.LEFT
        img = Image.open(io.BytesIO(png))
        w = min(width, height * img.width / img.height)
        x = self.LEFT + (width - w) / 2
        self.page.insert_image(pymupdf.Rect(x, self.y, x + w, self.y + height), stream=png)
        self.y += height + 6
        self.para(caption, size=9, fontname="hebo", gap=10)

    def footnotes(self, notes: list[str]):
        y = self.BOTTOM - 4 - 9 * len(notes)
        self.page.draw_line((self.LEFT, y - 4), (self.LEFT + 150, y - 4), width=0.5)
        for i, n in enumerate(notes):
            self.page.insert_text((self.LEFT, y + 7 + i * 9), n, fontname="helv", fontsize=6.8)

    def save(self, path: Path):
        total = len(self.doc)
        for n, page in enumerate(self.doc, start=1):   # running header + footer on every page
            page.insert_text((self.LEFT, 36), self.header, fontname="helv", fontsize=7.5, color=(0.35, 0.35, 0.35))
            page.draw_line((self.LEFT, 42), (self.RIGHT, 42), width=0.4, color=(0.6, 0.6, 0.6))
            page.draw_line((self.LEFT, 752), (self.RIGHT, 752), width=0.4, color=(0.6, 0.6, 0.6))
            page.insert_text((self.LEFT, 764), f"{DISCLAIMER}  |  Page {n} of {total}", fontname="helv", fontsize=7.5, color=(0.35, 0.35, 0.35))
        self.doc.save(str(path), garbage=4, deflate=True, deflate_images=True)


def write_guideline_pdf(path: Path) -> None:
    pdf = PdfWriter("Heart Failure with Reduced Ejection Fraction: Guideline Summary 2026",
                    "Synthetic Clinical Guideline Summary  |  HFrEF  |  CARD-GL-2026 Rev 3")
    p = pdf.page
    p.insert_text((pdf.LEFT, pdf.y + 18), "Heart Failure with Reduced Ejection Fraction (HFrEF):", fontname="hebo", fontsize=18)
    p.insert_text((pdf.LEFT, pdf.y + 40), "Guideline Summary 2026", fontname="hebo", fontsize=18)
    pdf.y += 56
    pdf.para("Document ID: CARD-GL-2026 | Revision: 3 | Classification: Public | Owner: Synthetic Guideline Consortium", size=8.5)
    pdf.para("Summary. This synthetic summary condenses the current standard of care for adults with heart failure with reduced "
             "ejection fraction (HFrEF): diagnosis, guideline-directed medical therapy (GDMT), device therapy, advanced care, "
             "emerging therapies and quality measures. It was written to demonstrate document extraction and is not clinical guidance.",
             size=10)
    pdf.columns(
        [("h", "1. Definition and Epidemiology"),
         ("p", ("HFrEF is heart failure with a left ventricular ejection fraction (LVEF) of 40% or lower. It accounts for roughly half "
               "of heart failure cases in this summary's reference population. Patients with an LVEF between 41% and 49% are classified "
               "as mildly reduced ejection fraction and are managed separately.")),
         ("p", ("Five-year mortality after a first heart failure hospitalization remains close to 50% in the reference registry, and "
               "30-day readmission after discharge averages 21%, which makes HFrEF a priority for value-based care programs."))],
        [("h", "2. Diagnosis"),
         ("p", ("Diagnosis combines symptoms, natriuretic peptides and imaging. An NT-proBNP above 125 pg/mL in the outpatient setting "
               "supports the diagnosis; transthoracic echocardiography confirms the LVEF and identifies structural causes.")),
         ("p", ("Initial work-up includes ECG, renal function, potassium, iron studies and screening for ischemic heart disease. "
               "Iron deficiency (ferritin below 100 ng/mL, or 100-299 ng/mL with transferrin saturation below 20%) should be corrected."))])
    pdf.heading("3. Guideline-Directed Medical Therapy (GDMT)")
    pdf.columns(
        [("p", ("The foundation of HFrEF care is four drug classes started together rather than sequentially: an angiotensin "
               "receptor-neprilysin inhibitor (ARNI), an evidence-based beta-blocker, a mineralocorticoid receptor antagonist (MRA) "
               "and an SGLT2 inhibitor. Each class independently reduces mortality or hospitalization."))],
        [("p", ("Rapid sequencing - all four pillars within six weeks and dose titration every two weeks - is preferred over "
               "waiting for target doses of one drug before adding the next. Table 1 lists recommended agents and target doses."))])
    pdf.table("Table 1. Recommended GDMT for HFrEF",
              ["Drug class", "Example agents", "Target dose (example)", "Recommendation", "Evidence"],
              [["ARNI", "sacubitril/valsartan", "97/103 mg twice daily", "Class 1; preferred over ACEi/ARB in NYHA II-III", "A"],
               ["ACE inhibitor", "enalapril", "10-20 mg twice daily", "Class 1 when ARNI is not feasible", "A"],
               ["ARB", "valsartan", "160 mg twice daily", "Class 1 if ACE inhibitor intolerant", "A"],
               ["Beta-blocker", "carvedilol, metoprolol succinate, bisoprolol", "carvedilol 25 mg twice daily", "Class 1 for all patients with HFrEF", "A"],
               ["MRA", "spironolactone, eplerenone", "spironolactone 25-50 mg daily", "Class 1 if eGFR above 30 and potassium below 5.0", "A"],
               ["SGLT2 inhibitor", "dapagliflozin, empagliflozin", "10 mg daily", "Class 1 with or without diabetes", "A"],
               ["Loop diuretic", "furosemide, torsemide", "titrate to euvolemia", "Class 1 for congestion", "B"],
               ["Hydralazine + isosorbide dinitrate", "fixed-dose combination", "75/40 mg three times daily", "Class 1 for self-identified Black patients, NYHA III-IV", "A"],
               ["Ivabradine", "ivabradine", "7.5 mg twice daily", "Class 2a; sinus rhythm with heart rate 70 or higher", "B"],
               ["sGC stimulator", "vericiguat", "10 mg daily", "Class 2b after a worsening HF event", "B"],
               ["IV iron", "ferric carboxymaltose", "weight-based dosing", "Class 2a for iron deficiency", "B"]],
              [100, 110, 105, 145, 52])
    pdf.heading("4. Device Therapy and Advanced Care")
    pdf.columns(
        [("p", ("Device decisions follow at least three months of optimized GDMT, because LVEF often improves with therapy. "
               "Figure 1 summarizes the treatment algorithm, including the thresholds for an implantable cardioverter-defibrillator "
               "and cardiac resynchronization therapy."))],
        [("p", ("Patients who remain in NYHA class IV despite therapy (Stage D) should be referred early to an advanced heart failure "
               "program. Palliative care is integrated alongside life-prolonging treatment, not only at the end of life."))])
    pdf.figure(algorithm_png(), 300, "Figure 1. HFrEF treatment algorithm: GDMT initiation, reassessment and device or advanced-therapy decisions.")
    pdf.heading("5. Emerging Therapies")
    pdf.columns(
        [("p", ("Development is concentrated in four areas: cardiac myosin activation to improve contractility, nonsteroidal "
               "mineralocorticoid receptor antagonists with lower hyperkalemia risk, gene therapy targeting calcium handling "
               "(SERCA2a), and device-based monitoring that detects congestion before symptoms.")),
         ("p", ("Digital and care-delivery programs that close the GDMT gap are a fast-growing category, because fewer than one "
               "in four eligible patients receives all four pillars at target dose."))],
        [("p", ("Figure 2 shows the active synthetic pipeline by modality and the share in late-stage development. Trial-level "
               "detail (sponsors, lead sites, endpoints and timelines) is maintained in the companion registry file.")),
         ("p", ("Strategy teams should track readouts expected in 2027-2029 for the late-stage programs, which could change "
               "first-line combinations and device referral volumes."))])
    pdf.figure(pipeline_chart_png("HFrEF"), 220, "Figure 2. Active HFrEF pipeline by modality, synthetic registry 2026.")
    pdf.footnotes(["a. Doses are examples for orientation only; prescribing follows the full guideline and product labeling.",
                   "b. Pipeline counts exclude completed studies. Synthetic registry snapshot dated 2026-09-30."])
    pdf.heading("6. Quality Measures and Care Delivery")
    pdf.para("Health systems are benchmarked on the measures below. Targets are synthetic values for demonstration.", size=10)
    pdf.table("Table 2. HFrEF quality measures",
              ["Measure", "Definition", "Benchmark target"],
              [["GDMT four-pillar rate", "Eligible patients on ARNI/ACEi/ARB, beta-blocker, MRA and SGLT2i", "80% or higher"],
               ["7-day follow-up", "Clinic or virtual visit within 7 days of discharge", "60% or higher"],
               ["30-day readmission", "All-cause readmission within 30 days of an HF discharge", "below 20%"],
               ["Cardiac rehabilitation referral", "Referral at discharge for eligible patients", "70% or higher"],
               ["Device evaluation", "ICD/CRT assessment documented after 3 months of GDMT", "90% or higher"]],
              [130, 252, 130])
    pdf.heading("7. Sources")
    pdf.para("This document is a synthetic composite prepared for a portfolio demonstration. It paraphrases widely published "
             "standard-of-care principles; all figures, targets and registry data are fictional.", size=9)
    pdf.save(path)


# ---------------------------------------------------------------- HTML guideline (second condition)
def write_diabetes_html(path: Path) -> None:
    path.write_text("""<!doctype html><html><head><title>Type 2 Diabetes: Standards of Care Summary 2026</title></head><body>
<h1>Type 2 Diabetes: Standards of Care Summary 2026</h1>
<p>Document ID: DM-GL-2026 | Revision: 2 | Classification: Public | Owner: Synthetic Guideline Consortium</p>
<p>This synthetic summary condenses the current standard of care for adults with type 2 diabetes. It is written for a document-AI
demonstration and is not clinical guidance.</p>
<h2>1. Diagnosis and Targets</h2>
<p>Type 2 diabetes is diagnosed with an A1C of 6.5% or higher, a fasting plasma glucose of 126 mg/dL or higher, or a 2-hour glucose of
200 mg/dL or higher during an oral glucose tolerance test. An A1C target below 7% is appropriate for many non-pregnant adults; less
stringent targets (below 8%) suit people with limited life expectancy or frequent hypoglycemia.</p>
<h2>2. First-Line Management</h2>
<p>Lifestyle therapy - medical nutrition therapy, at least 150 minutes per week of physical activity and diabetes self-management
education - is the foundation. Metformin remains a common first-line agent when there are no compelling cardiorenal indications.</p>
<h2>3. Cardiorenal Risk-Based Therapy</h2>
<p>For people with established atherosclerotic cardiovascular disease, heart failure or chronic kidney disease, an SGLT2 inhibitor or a
GLP-1 receptor agonist with proven benefit is recommended independent of A1C and of metformin use. Heart failure and CKD favor an SGLT2
inhibitor; weight management goals favor a GLP-1 receptor agonist or a dual GIP/GLP-1 receptor agonist.</p>
<table>
<tr><th>Drug class</th><th>A1C lowering</th><th>Weight effect</th><th>Cardiovascular / kidney benefit</th><th>Hypoglycemia risk</th></tr>
<tr><td>Metformin</td><td>High</td><td>Neutral</td><td>Potential benefit</td><td>Low</td></tr>
<tr><td>SGLT2 inhibitor</td><td>Intermediate</td><td>Loss</td><td>HF and CKD benefit</td><td>Low</td></tr>
<tr><td>GLP-1 receptor agonist</td><td>High</td><td>Loss</td><td>ASCVD benefit (selected agents)</td><td>Low</td></tr>
<tr><td>Dual GIP/GLP-1 receptor agonist</td><td>Very high</td><td>Large loss</td><td>Under study</td><td>Low</td></tr>
<tr><td>DPP-4 inhibitor</td><td>Intermediate</td><td>Neutral</td><td>Neutral</td><td>Low</td></tr>
<tr><td>Basal insulin</td><td>Highest</td><td>Gain</td><td>Neutral</td><td>High</td></tr>
<tr><td>Sulfonylurea</td><td>High</td><td>Gain</td><td>Neutral</td><td>High</td></tr>
</table>
<h2>4. Technology and Monitoring</h2>
<p>Continuous glucose monitoring (CGM) is recommended for adults on insulin and is increasingly used for people on non-insulin therapy.
Time in range (70-180 mg/dL) above 70% is the common CGM target. A1C is checked at least twice a year when at target.</p>
<h2>5. Emerging Therapies</h2>
<p>Pipeline activity centers on oral and multi-agonist incretin therapies, once-weekly basal insulin, glucokinase activators,
long-wear CGM devices and beta-cell replacement. Remission through structured weight loss is an expanding care-delivery model. See
the synthetic registry for programs and sponsors.</p>
<h2>6. Quality Measures</h2>
<ul><li>A1C control below 8%: target 70% of the diabetes panel or higher.</li>
<li>Annual kidney health evaluation (eGFR and urine albumin-to-creatinine ratio): target 80% or higher.</li>
<li>Statin therapy for people aged 40-75: target 85% or higher.</li></ul>
</body></html>
""", encoding="utf-8")


# ---------------------------------------------------------------- Word strategy brief (internal)
def write_strategy_docx(path: Path) -> None:
    d = docx.Document()
    d.core_properties.title = "Cardiometabolic Service Line Strategy Brief 2026"
    d.add_heading("Cardiometabolic Service Line Strategy Brief 2026", 0)
    d.add_paragraph("Document ID: STRAT-2026 | Revision: 1 | Classification: Internal | Owner: Riverbend Health Strategy Office")
    d.add_heading("1. Heart Failure Program Position", 1)
    d.add_paragraph("Riverbend Health admitted 2,230 heart failure patients in the last four quarters. The GDMT four-pillar rate at "
                    "discharge is 41%, well below the 80% benchmark, and the 30-day readmission rate is 21.4% after a peak of 24.8%. "
                    "Competitor St. Brendan Regional opened a dedicated HF titration clinic in 2025 and is recruiting for two "
                    "industry trials.")
    d.add_picture(io.BytesIO(admissions_chart_png()), width=Inches(6))
    d.add_heading("2. Partnership and Trial-Site Opportunities", 1)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, ["Partner", "Opportunity", "Fit", "Next step"]):
        cell.text = text
    for row in [["Cardiora Therapeutics", "Site for SYN-HF-101 Phase 3 expansion", "High - HF clinic volume", "Feasibility call in Q4"],
                ["Pentara Health AI", "TitrateAI pilot to lift the GDMT rate", "High - closes the 41% vs 80% gap", "Data-sharing review"],
                ["Meridian Cardio Devices", "Remote monitoring program (MER-Sense)", "Medium - needs device clinic staffing", "Business case"],
                ["Glycova Bio", "Site for SYN-DM-208 outcomes study", "Medium - large diabetes panel", "Await site selection"]]:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    d.add_heading("3. Diabetes Program Position", 1)
    d.add_paragraph("The primary-care diabetes panel is 38,400 adults; 58% have an A1C below 8%, short of the 70% target. CGM use "
                    "outside insulin users is 12%. A remission pathway and pharmacist-led titration are the lowest-cost levers.")
    d.add_heading("4. Recommendations for Leadership", 1)
    for rec in ["Fund a pharmacist-led GDMT titration clinic in FY27 (target four-pillar rate of 65% within 12 months).",
                "Apply to be a site for two late-stage HF trials to strengthen academic positioning.",
                "Pilot CGM-linked coaching for non-insulin patients with an A1C above 9%."]:
        d.add_paragraph(rec, style="List Bullet")
    d.save(str(path))


def write_trials_csv(path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(FIELDS)
        w.writerows(TRIALS)


if __name__ == "__main__":
    RAW.mkdir(parents=True, exist_ok=True)
    write_guideline_pdf(RAW / "CARD-GL-2026_hfref_guideline_summary.pdf")
    write_diabetes_html(RAW / "DM-GL-2026_type2_diabetes_care_summary.html")
    write_strategy_docx(RAW / "STRAT-2026_cardiometabolic_service_line.docx")
    write_trials_csv(RAW / "pipeline_trials.csv")
    for p in sorted(RAW.iterdir()):
        print(f"{p.name:50s} {p.stat().st_size / 1024:7.1f} KB")
