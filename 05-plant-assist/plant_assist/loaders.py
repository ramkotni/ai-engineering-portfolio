"""Extraction: turn source files (Markdown, Word, PDF) into one normalised form - Markdown text
plus document metadata - so a single structure-aware chunker handles every format.

  .md / .txt  YAML front matter (doc_id, title, revision, owner) + Markdown body
  .docx       heading styles -> #/##/###, paragraphs, list items, tables -> Markdown tables (body order kept)
  .pdf        page text + tables (pdfplumber) -> Markdown tables; metadata from file name / first line

Tables are always emitted as Markdown pipe tables so the chunker can keep them intact.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

import yaml

SUPPORTED_SUFFIXES = (".md", ".txt", ".docx", ".pdf")


@dataclass
class LoadedDocument:
    path: Path
    markdown: str
    meta: dict = field(default_factory=dict)

    @property
    def doc_id(self) -> str:
        return str(self.meta.get("doc_id") or self.path.stem)

    @property
    def title(self) -> str:
        return str(self.meta.get("title") or self.path.stem)


def load_document(path: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return _load_markdown(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pdf":
        return _load_pdf(path)
    raise ValueError(f"Unsupported file type: {path}")


# ---------------------------------------------------------------- Markdown
def parse_front_matter(text: str) -> tuple[dict, str]:
    """Split a '---' delimited YAML header from the body."""
    if text.startswith("---"):
        _, header, body = text.split("---", 2)
        return yaml.safe_load(header) or {}, body.lstrip()
    return {}, text


def _load_markdown(path: Path) -> LoadedDocument:
    meta, body = parse_front_matter(path.read_text(encoding="utf-8"))
    return LoadedDocument(path, body, meta)


# ---------------------------------------------------------------- tables
def table_to_markdown(rows: list[list[str | None]]) -> str:
    """Render extracted table rows as a Markdown pipe table (first row = header)."""
    clean = [[" ".join(str(c or "").split()).replace("|", "/") for c in row] for row in rows if row]
    clean = [r for r in clean if any(r)]
    if not clean:
        return ""
    width = max(len(r) for r in clean)
    clean = [r + [""] * (width - len(r)) for r in clean]
    lines = ["| " + " | ".join(clean[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in clean[1:]]
    return "\n".join(lines)


# ---------------------------------------------------------------- Word
_DOCX_KEYS = {"document id": "doc_id", "doc id": "doc_id", "revision": "revision", "owner": "owner", "effective": "effective"}


def _load_docx(path: Path) -> LoadedDocument:
    import docx  # python-docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    meta: dict = {}
    if document.core_properties.title:
        meta["title"] = document.core_properties.title

    parts: list[str] = []
    for child in document.element.body.iterchildren():  # keeps paragraphs and tables in reading order
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            if style == "title" or style.startswith("heading"):
                level = 1 if style == "title" else min(int(re.sub(r"\D", "", style) or 1), 3)
                parts.append(f"{'#' * level} {text}")
                if level == 1:
                    meta.setdefault("title", text)
            elif "list" in style:
                parts.append(f"- {text}")
            else:
                key, _, value = text.partition(":")
                if key.strip().lower() in _DOCX_KEYS and value.strip() and len(parts) < 4:
                    meta[_DOCX_KEYS[key.strip().lower()]] = value.strip()  # header block like "Document ID: X"
                else:
                    parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            parts.append(table_to_markdown([[cell.text for cell in row.cells] for row in table.rows]))
    return LoadedDocument(path, "\n\n".join(p for p in parts if p), meta)


# ---------------------------------------------------------------- PDF
def _load_pdf(path: Path) -> LoadedDocument:
    """PDFs carry no structure tags, so structure is reconstructed:
    - reading order: each page is cut into horizontal strips between tables, so text and
      tables stay in the order they appear;
    - headings: lines whose font is clearly larger than the body font become #/##/###;
    - metadata: "Document ID: ..." style lines near the top, as for Word files."""
    import statistics

    import pdfplumber

    meta: dict = {}
    segments: list[tuple[str, object]] = []  # ("lines", [(text, size)]) | ("table", markdown)
    with pdfplumber.open(str(path)) as pdf:
        if (pdf.metadata or {}).get("Title"):
            meta["title"] = pdf.metadata["Title"]
        for page in pdf.pages:
            y = 0.0
            for table in sorted(page.find_tables(), key=lambda t: t.bbox[1]):
                segments.append(("lines", _pdf_lines(page.crop((0, y, page.width, table.bbox[1])))))
                segments.append(("table", table_to_markdown(table.extract())))
                y = table.bbox[3]
            segments.append(("lines", _pdf_lines(page.crop((0, y, page.width, page.height)))))

    sizes = [size for kind, lines in segments if kind == "lines" for _, size in lines]
    body = statistics.median(sizes) if sizes else 0
    heading_sizes = sorted({round(s, 1) for s in sizes if s >= body * 1.15}, reverse=True)[:3]

    parts: list[str] = []
    for kind, content in segments:
        if kind == "table":
            parts.append(content)
            continue
        for text, size in content:
            if round(size, 1) in heading_sizes:
                level = heading_sizes.index(round(size, 1)) + 1
                parts.append(f"{'#' * level} {text}")
                if level == 1:
                    meta.setdefault("title", text)
                continue
            key, _, value = text.partition(":")
            if key.strip().lower() in _DOCX_KEYS and value.strip() and len(parts) < 4:
                meta[_DOCX_KEYS[key.strip().lower()]] = value.strip()
            elif parts and not parts[-1].startswith(("#", "|")) and not text[:1].isupper():
                parts[-1] += " " + text  # continuation of a wrapped line
            else:
                parts.append(text)
    markdown = "\n\n".join(p for p in parts if p)
    meta.setdefault("title", (markdown.splitlines() or [path.stem])[0].lstrip("# ")[:120])
    return LoadedDocument(path, markdown, meta)


def _pdf_lines(region) -> list[tuple[str, float]]:
    """Text lines in a page region with their (max) font size."""
    out = []
    for line in region.extract_text_lines(return_chars=True):
        text = " ".join(line["text"].split())
        if text:
            out.append((text, max(c["size"] for c in line["chars"])))
    return out
