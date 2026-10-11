"""Generates the synthetic Word document used to demonstrate DOCX extraction (headings + tables).

    python data/make_sample_docs.py

Output: data/docs/maintenance/MM-HP-200-TS_torque_specs.docx
"""

from pathlib import Path

import docx

OUT = Path(__file__).parent / "docs" / "maintenance" / "MM-HP-200-TS_torque_specs.docx"

TORQUE_SPECS = [
    ("Fastener", "Location", "Size", "Torque (ft-lb)", "Pattern / note"),
    ("Gland retainer bolts", "Main cylinder", "M20", "210", "Star pattern, 2 passes"),
    ("Cylinder mounting bolts", "Crown", "M30", "650", "Star pattern, 3 passes"),
    ("Platen clamp bolts", "Moving platen", "M24", "380", "Cross pattern"),
    ("Guide post nuts", "Columns", "M64", "1,450", "Hydraulic tensioner only"),
    ("Pump coupling set screws", "Power unit", "M10", "35", "Blue thread locker"),
    ("Accumulator mounting U-bolts", "Power unit", "1/2 in", "60", "Re-check after 24 h"),
]

FLUIDS = [
    ("Fluid", "Specification", "Capacity", "Change interval"),
    ("Hydraulic oil", "ISO VG 46 anti-wear", "120 gal", "4,000 h or annually"),
    ("Guide post grease", "EP-2 lithium", "-", "Monthly"),
    ("Accumulator gas", "Dry nitrogen", "1,400 psi pre-charge", "Check monthly"),
]


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value


def main() -> None:
    d = docx.Document()
    d.core_properties.title = "HP-200 Torque and Fluid Specifications"
    d.add_heading("HP-200 Torque and Fluid Specifications", level=1)
    d.add_paragraph("Document ID: MM-HP-200-TS")
    d.add_paragraph("Revision: 2")
    d.add_paragraph("Owner: Maintenance Planner")

    d.add_heading("Torque Specifications", level=2)
    d.add_paragraph(
        "Use a calibrated torque wrench (see QCS-004). Lubricate threads with clean ISO VG 46 oil "
        "unless the note says otherwise. Apply LOTO per SOP-SAF-001 before working on the press."
    )
    add_table(d, TORQUE_SPECS)

    d.add_heading("Fluid Specifications", level=2)
    add_table(d, FLUIDS)

    d.add_heading("Escalation", level=2)
    d.add_paragraph("Contact the Maintenance Planner (ext. 4410) for any fastener not listed here.")
    d.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
