"""Shared slide helpers (palette, shapes, tables, slide template) for the decks in docs/.

Importing this module creates the Presentation object `prs`; each deck script
adds slides to it and saves it. Set `deck_common.FOOTER` to change the footer.
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "eval" / "reports" / "latest.json"

GREEN = RGBColor(0x1F, 0x7A, 0x5A)
DARK = RGBColor(0x1D, 0x2B, 0x26)
MUTED = RGBColor(0x5F, 0x6F, 0x69)
LIGHT = RGBColor(0xEE, 0xF5, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xC2, 0x6B, 0x1E)
BLUE = RGBColor(0x2D, 0x5B, 0x9A)
RED = RGBColor(0xA8, 0x3A, 0x32)
GREY = RGBColor(0xF3, 0xF4, 0xF4)
FOOTER = "Evergreen Member Assist  ·  Agentic RAG on LangGraph"  # override per deck

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


# ------------------------------------------------------------- helpers ----
def text_box(slide, x, y, w, h, text, size=16, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
        r.font.name = "Segoe UI"
    return tb


def bullets(slide, x, y, w, h, items, size=17, color=DARK, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        level = 1 if item.startswith("  ") else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = ("–  " if level else "•  ") + item.strip()
        r.font.size = Pt(size - 2 * level)
        r.font.color.rgb = MUTED if level else color
        r.font.name = "Segoe UI"
        if level:
            p.level = 1
    return tb


def box(slide, x, y, w, h, text, fill=LIGHT, line=GREEN, color=DARK, size=13, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line_txt in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line_txt
        r.font.size = Pt(size if i == 0 else size - 2)
        r.font.bold = bold if i == 0 else False
        r.font.color.rgb = color
        r.font.name = "Segoe UI"
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED, label=None, dashed=False, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if dashed:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if label:
        mx, my = (x1 + x2) // 2, (y1 + y2) // 2
        text_box(slide, Emu(mx - Inches(0.8)), Emu(my - Inches(0.32)), Inches(1.6), Inches(0.3), label, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    return c


def polyline(slide, pts, color=MUTED, label=None, label_at=None):
    """Elbow path through points (inches); arrowhead on the last segment."""
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        arrow(slide, Inches(x1), Inches(y1), Inches(x2), Inches(y2), color=color, head=i == len(pts) - 2)
    if label:
        lx, ly = label_at
        text_box(slide, Inches(lx), Inches(ly), Inches(2.6), Inches(0.3), label, size=10, color=color, align=PP_ALIGN.CENTER)


def table(slide, x, y, w, rows, col_widths=None, size=12, header_fill=GREEN, row_h=0.42):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows))).table
    if col_widths:
        for i, cw in enumerate(col_widths):
            t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else GREY)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.name = "Segoe UI"
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else DARK
    return t


def new_slide(title, subtitle=None, notes=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()
    text_box(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.7), title, size=30, bold=True)
    if subtitle:
        text_box(s, Inches(0.6), Inches(1.0), Inches(12), Inches(0.45), subtitle, size=15, color=MUTED)
    footer = f"{FOOTER}  ·  {len(prs.slides)}"
    text_box(s, Inches(0.6), Inches(7.0), Inches(12), Inches(0.3), footer, size=10, color=MUTED)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def load_eval(path: Path = REPORT) -> dict | None:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return None


MONO = "Consolas"
CODE_BG = RGBColor(0x1E, 0x26, 0x23)
CODE_FG = RGBColor(0xE6, 0xEE, 0xEA)


def code_box(slide, x, y, w, h, code, size=11, title=None):
    if title:
        text_box(slide, x, y - Inches(0.36), w, Inches(0.34), title, size=13, bold=True, color=GREEN)
    n_lines = len(code.strip("\n").split("\n"))
    h = min(h, Inches(0.3 + n_lines * size * 1.25 / 72))  # shrink to fit content
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = CODE_BG
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    for i, line in enumerate(code.strip("\n").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = line if line else " "
        r.font.name = MONO
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0x8F, 0xD1, 0xB0) if line.lstrip().startswith("#") else CODE_FG
    return s
